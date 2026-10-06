from pathlib import Path
import json
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
TASKS=ROOT/'Practical Task'
EVIDENCE=ROOT/'evidence'/'validation'
NAME='2638120_Thaslim_Mohammed_Sahl_High_Performance_Computing'
FONT='Times New Roman'

def set_font(font,element,size=None):
    font.name=FONT
    if size is not None: font.size=Pt(size)
    properties=element.get_or_add_rPr()
    fonts=properties.find(qn('w:rFonts'))
    if fonts is None:
        fonts=OxmlElement('w:rFonts'); properties.insert(0,fonts)
    for script in ('ascii','hAnsi','eastAsia','cs'):
        fonts.set(qn('w:'+script),FONT)
    for attribute in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        fonts.attrib.pop(qn('w:'+attribute),None)

def index_field(doc, instruction):
    paragraph=doc.add_paragraph()
    for kind in ('begin','instruction','separate','end'):
        run=paragraph.add_run()
        if kind=='instruction':
            field=OxmlElement('w:instrText')
            field.set(qn('xml:space'),'preserve')
            field.text=instruction
        else:
            field=OxmlElement('w:fldChar')
            field.set(qn('w:fldCharType'),kind)
            if kind=='begin': field.set(qn('w:dirty'),'true')
        run._r.append(field)

def table_of_contents(doc):
    doc.add_page_break()
    doc.add_paragraph('Table of contents','TOC Heading')
    index_field(doc,' TOC \\o "1-2" \\h \\z \\u ')
    settings=doc.settings.element
    update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true')
    settings.append(update)
    doc.add_page_break()
    doc.add_paragraph('Table of figures','TOC Heading')
    index_field(doc,' TOC \\h \\z \\c "Figure" ')
    doc.add_page_break()

def figure_caption(doc, caption):
    match=re.fullmatch(r'Figure (\d+)\. (.+)',caption)
    assert match, caption
    paragraph=doc.add_paragraph(style='Caption')
    paragraph.add_run('Figure ')
    for kind in ('begin','instruction','separate','number','end'):
        run=paragraph.add_run()
        if kind=='instruction':
            field=OxmlElement('w:instrText')
            field.set(qn('xml:space'),'preserve')
            field.text=' SEQ Figure \\* ARABIC '
            run._r.append(field)
        elif kind=='number':
            run.text=match[1]
        else:
            field=OxmlElement('w:fldChar')
            field.set(qn('w:fldCharType'),kind)
            run._r.append(field)
    paragraph.add_run('. '+match[2])
    return paragraph

def link(doc,text,url):
    p=doc.add_paragraph()
    hyperlink=OxmlElement('w:hyperlink')
    rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    hyperlink.set(qn('r:id'),rid)
    run=OxmlElement('w:r'); properties=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts')
    for script in ('ascii','hAnsi','eastAsia','cs'): fonts.set(qn('w:'+script),FONT)
    properties.append(fonts)
    size=OxmlElement('w:sz'); size.set(qn('w:val'),'24'); properties.append(size)
    color=OxmlElement('w:color'); color.set(qn('w:val'),'164B6B'); properties.append(color)
    underline=OxmlElement('w:u'); underline.set(qn('w:val'),'single'); properties.append(underline)
    run.append(properties); t=OxmlElement('w:t'); t.text=text; run.append(t)
    hyperlink.append(run); p._p.append(hyperlink)
    return p

def code(doc,text):
    for line in text.splitlines():
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.space_before=Pt(0)
        run=p.add_run(line); set_font(run.font,run._element,10)

def section(doc,title,paragraphs):
    doc.add_heading(title,2)
    for paragraph in paragraphs: doc.add_paragraph(paragraph)

def colab_capture(doc,filename,caption,explanation,max_height=5.35):
    """Insert an actual browser capture without stretching or recreating it."""
    path=ROOT/'evidence'/filename
    assert path.is_file(), f'Actual Colab screenshot is required: {path}'
    with Image.open(path) as image:
        pixel_width,pixel_height=image.size
    available_width=(doc.sections[0].page_width-doc.sections[0].left_margin-
                     doc.sections[0].right_margin)/914400
    width=min(available_width,max_height*pixel_width/pixel_height)
    picture=doc.add_picture(str(path),width=Inches(width))
    picture._inline.docPr.set('descr',caption)
    paragraph=doc.paragraphs[-1]
    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next=True
    paragraph.paragraph_format.space_after=Pt(7)
    paragraph=figure_caption(doc,caption)
    paragraph.paragraph_format.keep_with_next=True
    doc.add_paragraph(explanation)

links=json.loads((ROOT/'colab_links.json').read_text()) if (ROOT/'colab_links.json').exists() else {}
task4_log=(EVIDENCE/'task4_public_run.log').read_text()
task4_times=re.search(r'images/download\.png:.*?CPU compute: ([\d.]+) ms; CUDA kernel: ([\d.]+) ms; CUDA transfers \+ kernel: ([\d.]+) ms', task4_log, re.S)
assert task4_times, 'Updated Task 04 execution log is required.'
task4_manifest=json.loads((EVIDENCE/'task4_validation.json').read_text())
assert task4_manifest['status']=='PASS'
task4_test_groups=len(task4_manifest['tests'])
task4_png_count=int(re.search(r'^(\d+) (?:inputs|PNGs)',task4_manifest['tests'][0]['details'])[1])
assert any('gradient axes' in test['test'] for test in task4_manifest['tests']), 'Four-view CUDA validation is required.'
task4_downloads=json.loads((TASKS/'Task 04'/'image_downloads.json').read_text())
assert len(task4_downloads)==1 and task4_downloads[0]['filename']=='download.png', 'Task 04 must demonstrate exactly one public input image.'
assert all(image['load_mode']=='live public download' for image in task4_downloads)
git_colab=lambda n:links.get(str(n),f'https://colab.research.google.com/github/sahl-2003/HPC/blob/main/Practical%20Task/Task%20{n:02d}/Task_{n:02d}.ipynb')

task_content={
1:{'title':'Word Occurrence Counting using Multithreading', 'source':'WordOccurrence.c',
   'command':'gcc -std=c11 -O2 -pthread WordOccurrence.c -o WordOccurrence\n./WordOccurrence WordOccurrenceDataset.txt 4',
   'objective':['Count every word in the command-line text file using Pthreads, merge the thread results, and write alphabetically ordered frequencies to result.txt. The user also supplies the thread count.'],
   'algorithm':['The host reads the file into one dynamically allocated buffer. A word is a run of ASCII letters or digits; letters are converted to lowercase. Spaces and punctuation separate words. For example, CUDA and cuda share one count, while don\'t is split into don and t. This rule is explicit so both the program and its independent checker interpret the same input.',
       'The initial byte slice is file_size / (8 * actual_threads), with a minimum of one byte. A mutex protects the shared nextByte index. A worker claims a slice, extends its end to the end of a word, updates nextByte, and releases the mutex. It then processes that distinct range without holding the lock. Workers request further slices until the file is exhausted.',
       'Each worker keeps its own growable occurrence array, sorts it with qsort, and compresses repeated words into local frequencies. Main joins all workers before combining their local counts, sorting the combined entries, and summing equal words. No worker writes the output file.'],
   'correctness':['Extending the end of each claimed slice prevents a word from being counted twice or split into fragments. The next slice starts exactly where the previous one ended. Every claimed byte range is disjoint and the final range ends at the file size.',
       'Only the work-claim index and allocation-failure flag need mutex protection. File bytes are read-only and count arrays belong to individual threads. pthread_join establishes that each local result is complete before the host reads it. The mutex is destroyed after all joins. Every word allocation, local array, merged array, thread array, and input buffer is released.'],
   'tests':['The supplied WordOccurrenceDataset.txt produced 120,000 words and 94 unique words. Results matched an independent Python Counter at 1, 2, 3, 7 and 16 threads, including their exact alphabetical order and total frequency.',
       'Boundary tests covered a 20,000-character word, CRLF, tabs, punctuation, digits, no final newline, and non-divisible workloads. An empty file produced a valid header with no word rows. Zero, negative, non-numeric, and out-of-range thread counts, missing files, and missing arguments returned errors.',
       'AddressSanitizer and UndefinedBehaviorSanitizer passed on the supplied dataset with seven threads. The program prints thread IDs and per-thread slice counts; their distribution can vary with scheduling without changing the final frequencies.'],
   'limits':['Counting and merging require memory proportional to the number of occurrences and local unique entries. Sorting takes O(W log W) across the collected words. File loading and the final merge are serial. The word definition is ASCII and does not perform Unicode language segmentation.',
       'The recorded single-run counting and merge times were 45.915 ms with one thread and 37.812 ms with two threads. These are observations from one run, not stable speedup estimates. Colab scheduling, allocation overhead, and serial work affect the times. More threads must not be assumed to improve performance.']},
2:{'title':'Multiple operations with Matrices using multithreading','source':'MatrixOperations.c',
   'command':'gcc -std=c11 -O2 -fopenmp MatrixOperations.c -lm -o MatrixOperations\n./MatrixOperations MatData.txt 4',
   'objective':['Read the input path from argv[1] and the requested OpenMP thread count from argv[2]. Process matrices as consecutive pairs and write the full set of valid results and inapplicable-operation messages to results.txt.'],
   'algorithm':['The input uses a rows,cols header followed by exactly rows lines of comma-separated numeric values. Blank lines are permitted between matrices. fopen, fscanf, fprintf and fclose perform file I/O. fscanf reads characters into a growable line buffer, so row boundaries remain significant. strtol and strtod validate headers and complete numeric tokens.',
       'Matrices use malloc for the row pointer array and each row of doubles, following the lecturer\'s matrix examples. The complete file is validated before results.txt is opened. Input errors identify the matrix, row or line, and allocated data is released. An odd matrix count is rejected because it cannot form complete pairs.',
       'For equal shapes the program computes A+B, A-B, A.*B and A./B. A zero divisor produces NaN in that cell. Both transposes are always computed. A*B is computed only when A.cols equals B.rows; its cell is the dot product sum over k. Every unavailable operation gets one clear message and the remaining operations continue.'],
   'correctness':['For each operation the result row count is the number of outer loop iterations. The thread cap is min(requested_threads, result_rows). This prevents more requested workers than result rows. The program prints the requested, capped and actual OpenMP team sizes.',
       'The outer row loop uses omp for with a static schedule. The input matrices are shared and read-only. Iteration variables and the dot-product sum are local to the calculation, and each output cell has one writer. The implicit barrier completes the computation before fprintf reads the result. A shared reduction is unnecessary because the sum belongs to one cell, not to the whole result.',
       'Every output includes the pair identifier, operation name and dimensions, followed by comma-separated rows. Results use 17 significant digits to preserve double precision. Row allocations are freed individually before the pointer array is freed. Overflow checks precede allocations.'],
   'tests':['All 50 supplied matrices were processed as 25 pairs. Every applicable output from requested thread counts of 1, 4 and 64 matched NumPy, using relative tolerance 1e-12 and absolute tolerance 1e-10. The checker also verified operation messages, pair order, dimensions and thread caps.',
       'Additional fixtures covered square matrices, rectangular matrix products, pairs with different shapes, one-row matrices and division by zero. Thirteen invalid files tested empty input, odd counts, bad headers, missing or extra row values, non-numeric tokens, NaN, infinity and numeric overflow. Missing files and invalid thread arguments were also rejected.',
       'AddressSanitizer and UndefinedBehaviorSanitizer passed on the supplied file with 64 requested threads. Pair 1 has two 3x4 matrices: element-wise operations and both transposes are valid, while the matrix product is not. Pair 2 has 4x6 and 6x2 matrices: the product and transposes are valid, while the four element-wise operations are not.'],
   'limits':['A matrix product of m x k and k x n requires O(mkn) arithmetic. Element-wise operations and transposes require O(mn). All input matrices remain in host memory until validation and processing finish; result memory is allocated one operation at a time.',
       'Many supplied matrices are small. Opening an OpenMP team for a short loop can cost more than its arithmetic. This implementation follows the required thread cap and records actual team sizes, but a small workload does not establish a speed advantage. Parsing, output formatting, and file writes remain serial.']},
3:{'title':'Password Cracking with Files using CUDA','source':'PWCrack.cu',
   'command':'nvcc -O2 -arch=sm_75 PWCrack.cu -o PWCrack\n./PWCrack passwords.txt',
   'objective':['Read every encrypted line from the command-line file, recover its two lowercase letters and two digits on the GPU, and write one plaintext answer per line to decrypted.txt in the original order.'],
   'algorithm':['The supplied CryptForCuda.c transformation creates ten encrypted characters from four raw characters. Its additions, subtractions, and boundary adjustments are retained exactly. The boundary adjustment is not replaced with conventional modulo wrapping. The device function writes into a caller-owned char[11] buffer, avoiding the shared static buffer used in the original serial example.',
       'The host validates each line as six lowercase letters followed by four digits. It allocates 11*N encrypted bytes, 5*N plaintext bytes, and N*sizeof(int) status bytes, and copies the encrypted array to device memory. The x thread index is blockIdx.x * blockDim.x + threadIdx.x; a boundary guard excludes padded threads.',
       'One CUDA thread owns one encrypted password. Nested loops generate the 26*26 letter pairs. The first six encrypted characters depend only on the letters, so a non-matching pair is rejected before trying the 100 digit pairs. A complete ten-character comparison determines success. This retains the lecturer\'s candidate-encryption logic while removing digit trials that cannot succeed.',
       'For 10,000 records the launch is 40 x blocks with 256 threads per block. Smaller inputs use a smaller block size and at least two x blocks. Larger inputs derive the number of blocks from the record count and device limit. Each thread writes only its own result and status, so no global success flag or atomic operation is needed.'],
   'correctness':['The CPU copies plaintext and status arrays back after checking the kernel launch and completion. Every recovered plaintext is re-encrypted on the host and compared with its original encrypted line before it is accepted. An uncrackable but syntactically valid record writes NOT_FOUND in its position and returns exit code 2.',
       'Input ordering is preserved by the fixed offsets record*11 and record*5. Every CUDA allocation, transfer, event, execution check and deallocation reports errors. Device buffers and events are released, and all host arrays are freed, including failure paths.'],
   'tests':['The normal dataset contains 10,000 synthetic passwords created with seed 6005 and the supplied encryption logic. It includes boundary cases such as aa00, zz99, az09 and za90. All 10,000 recovered values matched the known plaintext file exactly.',
       'An exhaustive independent test generated all 67,600 possible raw passwords, checked that their encrypted values were distinct within this domain, and recovered every one in order. Input sizes of 1, 2 and 257 tested launch padding and multiple blocks. A separate test compiled and ran the lecturer\'s PasswordGeneratorToText.c and recovered its 25 generated records.',
       'Empty files, short lines, long lines, uppercase characters, invalid digits and blank lines were rejected. The valid-shaped string aaaaaa0000 had no candidate and correctly produced NOT_FOUND. The recorded kernel time for 10,000 records was 0.303 ms on the Tesla T4; file reads, allocation, transfers and output writes are excluded from that event timing.'],
   'limits':['This is the assessment\'s synthetic character transformation, not an attack on a real authentication service. Candidate filtering is valid because encrypted positions 0..5 depend only on the two letters. Without this property, the same pruning would not be valid.',
       'All records are held in host and device memory at once. Memory demand therefore grows linearly with file size, and files that cannot be allocated fail clearly. Kernel time alone must not be reported as end-to-end application time or compared with a CPU program that includes file I/O.']},
4:{'title':'Sobel Edge Detection using Cuda across many PNG images','source':'SobelEdge.cu',
   'command':'nvcc -O2 -arch=sm_75 lodepng.cpp SobelEdge.cu -o SobelEdge\nmkdir -p outputs\n./SobelEdge outputs images/download.png',
   'objective':['Use one public input image, download.png, to demonstrate the brief\'s four views: Original Image, Gradient in X direction, Gradient in Y direction and Sobel Edge Detection. The four panels are views of the same input. Save both gradient maps and the combined edge map at the original dimensions. The CUDA program also accepts multiple PNG paths when a batch is required.'],
   'algorithm':['The notebook uses wget to download download.png (300x300) from the public image URL in the lecturer\'s last-class code and checks its SHA-256 checksum before use. The input needs no GitHub login or Drive mount. After lodepng.cpp and lodepng.h are uploaded, LodePNG decodes the file into an RGBA host array, with its original licence notices retained.',
       'The last-class sample demonstrates PNG decoding, explicit CUDA memory transfers and rgbToGray; its final kernel stops at grayscale conversion. This implementation extends that workflow with the assessed Sobel convolution. Luminance uses the lecturer\'s 0.30R + 0.59G + 0.11B weights as (30*R + 59*G + 11*B) / 100, rounded down. Integer arithmetic gives identical CPU and GPU truncation. Gx is [-1,0,1; -2,0,2; -1,0,1] and Gy is [-1,-2,-1; 0,0,0; 1,2,1]. Out-of-image neighbours are zero.',
       'Gx measures left-to-right intensity changes and Gy measures top-to-bottom changes, following the written axis definitions. The kernel saves min(255, abs(Gx)) and min(255, abs(Gy)) in separate one-byte-per-pixel display buffers. The final magnitude uses the original signed sums: sqrt(Gx*Gx + Gy*Gy), then truncation and saturation at 255. Clipping a display map never changes that calculation. The final edge value fills RGB and preserves input alpha; the separate gradient maps contain grayscale intensity only.',
       'RGBA input and final output arrays each use width*height*4 bytes. Each gradient array uses width*height bytes. The x grid has ceil(pixel_count/256) blocks and a guard for padded threads. The command accepts any number of PNGs. Output names are outImg_, outImg_Gx_ and outImg_Gy_ followed by the input basename. The host rejects duplicate or colliding generated names before processing. Gradient encoding uses lodepng_encode_file with LCT_GREY, 8, matching the lecturer\'s grayscale-buffer workflow.'],
   'correctness':['Checked kernel completion precedes device-to-host copies of all three maps. A serial C reference expresses the gradients as neighbour sums and compares every final RGBA byte and both gradient buffers. An independent NumPy checker verifies all saved maps and their dimensions, including alpha in the final edge image.',
       'Every successful output has exactly the original image dimensions. Zero padding can produce a visible edge around a constant bright image because its boundary neighbours are black. This is expected under the specified padding rule. PNG decode or encode failures and CUDA failures return errors; the batch can continue to a later valid input.',
       'Images are processed one at a time. Host image and reference buffers, all four device buffers and timing events are released after each image, including failure paths. Working memory depends on the largest image. Gradients are encoded losslessly from 8-bit grayscale buffers; LodePNG can optimise the file\'s storage without changing its decoded intensities.'],
   'tests':[f'The revised notebook downloaded its single demonstration image live on the Tesla T4. Every pixel in the X gradient, Y gradient and final edge maps matched the references across {task4_png_count} PNG test inputs. The public demonstration uses only download.png; additional small generated fixtures check zero padding, saturation, final-image transparency, grayscale rounding and padded launches.',
       'Directional fixtures independently check that a horizontal ramp has zero interior Gy and a vertical ramp has zero interior Gx. Negative gradients remain visible through their absolute values. A constructed 3x3 patch gives Gx=6, Gy=8 and final magnitude 10 at its centre, reproducing the brief\'s numerical example. Corrupt files, invalid paths, duplicate inputs and cross-output filename collisions are also tested; a batch continues to a valid image after a corrupt one.',
       f'The updated 300x300 download.png run took {task4_times[1]} ms for serial CPU computation, {task4_times[2]} ms for the CUDA kernel, and {task4_times[3]} ms for transfers plus kernel. These are single-run observations. GPU allocation, PNG decoding and encoding are outside the transfer-and-kernel interval. The first CUDA invocation has additional startup overhead, so timings are not stable benchmark estimates.'],
   'limits':['The implementation uses global memory and a basic one-thread-per-pixel kernel, matching the taught CUDA model. It does not use shared-memory tiling or concurrent image streams. Reading overlapping neighbourhoods repeats memory accesses, but the independent output ownership makes correctness straightforward.',
       'Converting colour to luminance detects intensity edges. It does not preserve coloured edge directions. Very small images can be slower on the GPU once transfer and launch costs are included. The report separates the timing intervals to avoid implying that kernel speed is total application speed.']}
}

capture_content={
1:('Figure 2. Task 01 word occurrence results in Google Colab.',
   'The captured program summary shows four requested and four actual Pthreads, 120,000 total words and 94 unique words. It names result.txt as the output file. The next cell displays part of the word-frequency file. These visible results document the supplied dataset run; the independent checks described above also test thread counts and boundary cases.'),
2:('Figure 4. Task 02 matrix operations in Google Colab.',
   'The visible summary reports 50 matrices processed as 25 pairs and names results.txt. The output preview below contains numeric matrix rows and a matrix multiplication heading with its result dimensions. This confirms that the program writes matrix results to the required file. The separate NumPy comparisons verify every applicable operation, including rows that are outside this screenshot.'),
3:('Figure 6. Task 03 CUDA password recovery in Google Colab.',
   'The program summary reports 10,000 recovered and verified passwords out of 10,000 and identifies decrypted.txt. The display cell shows encrypted records beside their recovered two-letter, two-digit plaintext values. The T4 runtime is visible in the Colab status bar. The full-domain and invalid-input tests described above provide checks beyond these displayed examples.'),
4:('Figure 8. Task 04 single-image CUDA execution in Google Colab.',
   'This capture records the notebook run for the 300x300 download.png image on the T4 runtime. Its output identifies the X gradient, Y gradient and final edge PNGs. The reference comparison checks these maps against the serial C calculation. The independent NumPy tests check the saved PNG pixels and gradient directions separately.')
}
four_views_caption='Figure 9. Four views of the same download.png input displayed in Google Colab.'
four_views_explanation=('The upper-left panel is the original colour image. The upper-right panel displays the absolute X gradient, which responds to left-right intensity changes; the lower-left displays the absolute Y gradient, which responds to top-bottom changes. The lower-right panel is the combined Sobel magnitude. Each map is calculated by the CUDA kernel from the same original pixels. Display values are clamped to 255, while the final magnitude uses the signed, unclipped Gx and Gy sums. All four views retain the 300x300 image dimensions.')

code_content={
1:{'caption':'Figure 1. Colab code for mutex-protected slice allocation in countWords.',
   'explanation':'I implemented dynamic work allocation with Pthreads. Each worker locks the shared cursor only while claiming its next slice, extends the slice to avoid splitting a word, then unlocks it before scanning the text. Word counts are stored in each worker\'s own collection before the results are merged.',
   'source':'WordOccurrence.c, countWords, lines 55 to 79'},
2:{'caption':'Figure 3. Colab code for parallel matrix rows and local product accumulation.',
   'explanation':'I implemented the matrix operations inside an OpenMP parallel region. Static scheduling gives different output rows to different threads, with the thread count capped by the number of rows. Each matrix-product cell uses its own local sum, so workers write separate result cells without sharing an accumulator.',
   'source':'MatrixOperations.c, operation, lines 188 to 210'},
3:{'caption':'Figure 5. Colab code for password indexing and candidate search.',
   'explanation':'I implemented a CUDA kernel in which each thread processes one encrypted password. The thread checks its index, uses its own candidate buffers and rejects letter pairs that fail to match the first six encrypted characters before trying digit pairs. A full match is copied to that thread\'s separate output slot.',
   'source':'PWCrack.cu, crackPasswords, lines 40 to 64'},
4:{'caption':'Figure 7. Colab code for pixel indexing, Sobel gradients and edge magnitude.',
   'explanation':'I implemented one CUDA thread per image pixel to calculate the X and Y gradients using the Sobel masks. Neighbours outside the image use zero padding, and the two signed gradient sums are combined with sqrt(Gx*Gx + Gy*Gy). I clamp the separate gradient display maps and the final magnitude to the 8-bit image range.',
   'source':'SobelEdge.cu, sobelEdge, lines 22 to 46'}
}

doc=Document()
for section_ in doc.sections:
    section_.top_margin=Inches(.8); section_.bottom_margin=Inches(.9)
    section_.footer_distance=Inches(.35)
    section_.left_margin=Inches(.85); section_.right_margin=Inches(.85)
styles=doc.styles
for style in styles:
    if hasattr(style,'font'): set_font(style.font,style._element)
for name in ['TOC Heading','TOC 1','TOC 2','Table of Figures']:
    if name not in styles: styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    set_font(styles[name].font,styles[name]._element)
styles['Normal'].font.size=Pt(12)
styles['Normal'].paragraph_format.space_after=Pt(7)
styles['Normal'].paragraph_format.line_spacing=1.1
for name in ['Title','Subtitle','TOC Heading']+[f'Heading {n}' for n in range(1,10)]:
    styles[name].font.color.rgb=RGBColor(0,0,0)
styles['Title'].font.size=Pt(26)
styles['Heading 1'].font.size=Pt(16)
styles['Heading 1'].paragraph_format.space_before=Pt(15)
styles['Heading 1'].paragraph_format.space_after=Pt(8)
styles['Heading 2'].font.size=Pt(13)
styles['Heading 2'].paragraph_format.space_before=Pt(12)
styles['Heading 2'].paragraph_format.space_after=Pt(6)
styles['Caption'].font.size=Pt(10)
styles['Caption'].font.color.rgb=RGBColor(0,0,0)
styles['Caption'].paragraph_format.space_after=Pt(7)
styles['TOC Heading'].font.size=Pt(18)
styles['TOC Heading'].font.bold=True
styles['TOC Heading'].paragraph_format.space_after=Pt(12)
for name in ['TOC 1','TOC 2']:
    styles[name].base_style=styles['Normal']
    styles[name].font.size=Pt(12)
    styles[name].paragraph_format.line_spacing=1
    styles[name].paragraph_format.space_after=Pt(0)
    styles[name].paragraph_format.space_before=Pt(3 if name=='TOC 1' else 0)
    styles[name].paragraph_format.left_indent=Inches(0 if name=='TOC 1' else .2)
styles['TOC 1'].font.bold=True
styles['Table of Figures'].base_style=styles['Normal']
styles['Table of Figures'].font.size=Pt(12)
styles['Table of Figures'].paragraph_format.space_after=Pt(8)
styles['Table of Figures'].paragraph_format.line_spacing=1.1
for border in list(styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
doc.core_properties.title='High Performance Computing portfolio'
doc.core_properties.author='Thaslim Mohammed Sahl'
doc.add_paragraph('High Performance Computing portfolio','Title')
doc.add_paragraph('6CS005 2025/26 assessment')
doc.add_paragraph('Thaslim Mohammed Sahl\nStudent number 2638120\n6 October 2026')
doc.add_paragraph('This portfolio implements the four tasks in the current assessment brief using Pthreads, OpenMP and CUDA. Each task has a separate source file, Colab notebook and output resources. Correctness is checked against independent reference calculations and invalid-input tests, with actual execution evidence from a Python 3 Colab runtime using a Tesla T4 GPU.')
table_of_contents(doc)
doc.add_heading('Execution and files',1)
doc.add_paragraph('The programs use the file-handling, dynamic allocation, thread structures, mutexes, parallel loops, CUDA index calculation, and explicit host-device transfer patterns from the supplied weeks and code sir.txt. Python performs setup, fixture generation, verification and display. The assessed computations are in C and CUDA; the CUDA sources remain separate .cu files.')
doc.add_paragraph('Task folders follow the Practical Task and Task 01 to Task 04 organisation. Output names used here are result.txt for word counts, results.txt for matrices, decrypted.txt for passwords, and outImg_ prefixed PNGs for edges. Task 04 adds outImg_Gx_ and outImg_Gy_ PNGs so all four illustrated views can be shown. No input resource requires mounting a private Google Drive folder.')
link(doc,'GitHub project','https://github.com/sahl-2003/HPC')
link(doc,'Download complete assignment and recordings','https://github.com/sahl-2003/HPC/releases/latest')
doc.add_paragraph('Open a task notebook in Google Colab and upload the files requested by its first cell: Task 01, WordOccurrenceDataset.txt; Task 02, MatData.txt; Task 03, passwords.txt and expected_passwords.txt; Task 04, lodepng.cpp and lodepng.h. Task 04 downloads download.png from its public URL. Run all cells to compile the program and display the output. Use a T4 GPU for the CUDA tasks. All four notebook links allow anyone with the link to view and copy the code. The GitHub project contains the source and resources.')
doc.add_heading('Validation result',1)
doc.add_paragraph(f'All 23 initial test groups passed on the Tesla T4, including multiple CPU thread counts, all supplied matrix pairs, exhaustive recovery of the 67,600-password domain, invalid inputs, AddressSanitizer and UndefinedBehaviorSanitizer. Task 04 was then rerun with separate X and Y maps: all {task4_test_groups} updated groups passed, including complete comparisons of three maps for each of {task4_png_count} PNG inputs. Raw logs and separate initial and updated validation manifests are saved under evidence/validation. Task 04 figures and timings below use the four-view version.')

for n,data in task_content.items():
    heading=doc.add_heading(f'Task {n:02d} {data["title"]}',1)
    heading.paragraph_format.page_break_before=True
    if n==4:
        heading.clear()
        heading.add_run('Task 04\n'+data['title'])
    link(doc,f'Task {n:02d} source notebook in Colab',git_colab(n))
    if str(n) in links: link(doc,f'Executed Task {n:02d} Colab notebook',links[str(n)])
    section(doc,'Objective',data['objective'])
    section(doc,'Implementation',data['algorithm'])
    doc.add_heading('Code I implemented',2)
    snippet=code_content[n]
    colab_capture(doc,f'Task_{n:02d}_Code_Colab.jpg',snippet['caption'],
                  snippet['explanation'],max_height=5.2)
    doc.add_paragraph('Source excerpt: '+snippet['source']+'. The full source remains in '+data['source']+'.')
    section(doc,'Synchronisation and memory',data['correctness'])
    doc.add_heading('Build and run',2); code(doc,data['command'])
    section(doc,'Results and tests',data['tests'])
    if n==2:
        code(doc,'Pair 1: A=3,4 B=3,4\nTranspose A - 4,3\nMatrix multiplication cannot be done (A.cols != B.rows).\nPair 2: A=4,6 B=6,2\nMatrix multiplication - 4,2')
        doc.add_paragraph('Selected output headings from the two different operation-applicability cases. Full numeric matrices are in results.txt.')
    elif n==3:
        log=(EVIDENCE/'task3_10000_passwords.log').read_text().splitlines()
        code(doc,'\n'.join(log[:9]))
        doc.add_paragraph('Actual CUDA execution excerpt. The full file has 10,000 recovered plaintext lines.')
    section(doc,'Performance and limits',data['limits'])
    heading=doc.add_heading('Colab execution evidence',2)
    heading.paragraph_format.page_break_before=True
    caption,explanation=capture_content[n]
    colab_capture(doc,f'Task_{n:02d}_Colab.jpg',caption,explanation)
    if n==4:
        heading=doc.add_heading('Task 04 image outputs',2)
        heading.paragraph_format.page_break_before=True
        colab_capture(doc,'Task_04_Four_Views_Colab.jpg',four_views_caption,
                      four_views_explanation,max_height=6.1)
    # A separate readable task answer is included in its task folder.
    lines=[f'# Task {n:02d} {data["title"]}','', 'Thaslim Mohammed Sahl | 2638120 | 6CS005','',
        f'[Open task notebook in Colab]({git_colab(n)})']
    if str(n) in links: lines+=['',f'[Executed Colab notebook]({links[str(n)]})']
    for title,key in [('Objective','objective'),('Implementation','algorithm'),('Synchronisation and memory','correctness'),('Results and tests','tests')]:
        lines+=['',f'## {title}','']
        for paragraph in data[key]: lines += [paragraph,'']
        if key=='algorithm':
            lines+=['## Code I implemented','',
                    f'![{code_content[n]["caption"]}](../../evidence/Task_{n:02d}_Code_Colab.jpg)','',
                    code_content[n]['explanation'],'',
                    'Source excerpt: '+code_content[n]['source']+'. The full source remains in '+data['source']+'.','']
    lines+=['','## Colab execution evidence','',
            f'![{capture_content[n][0]}](../../evidence/Task_{n:02d}_Colab.jpg)','',
            capture_content[n][1],'']
    if n==4:
        lines += [f'![{four_views_caption}](../../evidence/Task_04_Four_Views_Colab.jpg)',
                  '',four_views_explanation,'']
    lines += ['## Performance and limits','']
    for paragraph in data['limits']: lines += [paragraph,'']
    lines+=['## Build and run','','```sh',data['command'],'```']
    (TASKS/f'Task {n:02d}'/f'Task_{n:02d}_Report.md').write_text('\n'.join(lines),encoding='utf-8')

heading=doc.add_heading('Evidence and submission',1)
heading.paragraph_format.page_break_before=True
doc.add_paragraph('The evidence directory contains validation logs and four separate browser screen recordings named Task_01.mp4 to Task_04.mp4. Each recording captures the live Colab view while the corresponding program is rerun and its output is displayed. CUDA computation runs on Colab\'s GPU. The recordings capture the browser viewport at a reduced frame rate, with the original elapsed time preserved. Output files are saved beside each task source.')
doc.add_paragraph('The task sections contain actual Google Colab screenshots of the implementation code and its visible results. Each code excerpt has an explanation of the work I implemented, followed by its source location. Tasks 01 to 03 use output captures from their verified executions. Task 04 shows its current single-image execution and the four views of download.png. The screenshots accompany the saved outputs and validation logs; a partial output preview alone does not establish that every test passed.')
doc.add_paragraph('Each task is packaged separately with its C or CUDA source, notebook and resources. A complete portfolio archive also contains this report and all four task folders. Video evidence is kept locally in the evidence directory; large recording files are excluded from Git history.')
doc.add_heading('References',1)
for text in [
    'University of Wolverhampton (2025/26). 6CS005 Assessment 25-26. Supplied assignment brief.',
    '6CS005 teaching materials. Weeks 1 and 2: multithreading, Pthreads, mutexes and joining threads. Supplied lectures and workshops.',
    '6CS005 teaching materials. Weeks 3 and 4: C file handling, OpenMP parallel loops, shared and private values, and matrix operations. Supplied lectures and workshops.',
    '6CS005 teaching materials. Weeks 6 to 9: CUDA grids, thread indices, GPU memory, file input and PNG processing. Supplied lectures, workshops and source examples.',
    'Lecturer source collection. code sir.txt; CryptForCuda.c; PasswordGeneratorToText.c; OMPMatSumFromFile.c.txt; Negative.cu; last-class SobelEdge.cu PNG/grayscale example supplied as pasted text. Supplied teaching examples.',
    'Dissanayake, K. Portfolio reference. Consulted for folder organisation and report layout. The earlier tasks differ from the current brief.',
    'Vandevenne, L. (2018). LodePNG version 20180910. Unmodified PNG codec supplied in the reference resources; original licence retained.'
]: doc.add_paragraph(text)
link(doc,'NVIDIA CUDA Programming Guide','https://docs.nvidia.com/cuda/cuda-programming-guide/index.html')
link(doc,'OpenMP specification execution model','https://www.openmp.org/spec-html/5.2/openmpse3.html')
link(doc,'LodePNG documentation','https://lodev.org/lodepng/')
link(doc,'Lecturer PNG input used in Task 04','https://i.ibb.co/5gB80K0Z/download.png')
for paragraph in doc.paragraphs:
    if paragraph.style.name.startswith('Heading'):
        paragraph.paragraph_format.keep_with_next=True
footer=doc.sections[0].footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
set_font(styles['Footer'].font,styles['Footer']._element,10)
for target in (styles['Footer']._element,footer._p):
    for frame in list(target.iter(qn('w:framePr'))): frame.getparent().remove(frame)
footer.paragraph_format.left_indent=Inches(0)
footer.paragraph_format.right_indent=Inches(0)
footer.paragraph_format.first_line_indent=Inches(0)
footer.paragraph_format.tab_stops.clear_all()
footer.paragraph_format.line_spacing=1
footer.paragraph_format.space_before=Pt(0)
footer.paragraph_format.space_after=Pt(0)
styles['Footer'].paragraph_format.left_indent=Inches(0)
styles['Footer'].paragraph_format.right_indent=Inches(0)
styles['Footer'].paragraph_format.first_line_indent=Inches(0)
styles['Footer'].paragraph_format.tab_stops.clear_all()
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE')
run=footer.add_run('1'); set_font(run.font,run._element,10)
footer._p.remove(run._r); field.append(run._r); footer._p.append(field)
destination=ROOT/(NAME+'.docx')
doc.save(destination)
print(destination)
